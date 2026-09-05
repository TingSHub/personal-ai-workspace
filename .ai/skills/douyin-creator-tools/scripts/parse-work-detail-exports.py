#!/usr/bin/env python3
"""创作者中心单作品详情导出 Excel -> 复盘用结构化 JSON。"""

import json
import re
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

HEADER_MAP = {
    "封面点击率": "cover_click_rate", "平均播放时长": "avg_view_seconds", "完播率": "completion_rate",
    "2s跳出率": "bounce_rate_2s", "平均播放占比": "avg_view_proportion", "5s完播率": "completion_rate_5s",
    "时间": "time_range", "跳过率": "skip_rate", "回看率": "rewatch_rate",
    "点赞率": "like_rate", "评论率": "comment_rate", "分享率": "share_rate",
    "收藏率": "favorite_rate", "弹幕量": "bullet_count", "不感兴趣率": "not_interested_rate",
    "来源": "source", "来源占比": "source_share", "对比7日": "change_vs_7d",
    "涨粉量": "follower_gain", "涨粉率": "follower_gain_rate", "脱粉量": "follower_loss",
    "脱粉率": "follower_loss_rate", "不感兴趣量": "not_interested_count", "日期": "date",
    "抖音": "douyin_gain", "抖音精选": "douyin_select_gain",
}


def value(raw):
    if raw is None:
        return None
    if isinstance(raw, (int, float)):
        return raw
    text = str(raw).strip()
    if not text:
        return None
    if text.endswith("%"):
        try:
            return float(text[:-1])
        except ValueError:
            return text
    match = re.fullmatch(r"(-?\d+(?:\.\d+)?)秒", text)
    if match:
        return float(match.group(1))
    try:
        return float(text) if "." in text else int(text)
    except ValueError:
        return text


def read_workbook(path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheets = {}
    for sheet in workbook.worksheets:
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            continue
        headers = list(rows[0])
        records = []
        for row in rows[1:]:
            if not any(x is not None and str(x).strip() for x in row):
                continue
            record = {}
            for index, header in enumerate(headers):
                key = HEADER_MAP.get(str(header).strip()) if header is not None else None
                if key:
                    record[key] = value(row[index] if index < len(row) else None)
            if record:
                records.append(record)
        sheets[sheet.title] = {"headers": [str(x) if x is not None else None for x in headers], "rows": records}
    return sheets


def main():
    if len(sys.argv) != 3:
        print("用法: python3 parse-work-detail-exports.py <导出目录> <输出.json>", file=sys.stderr)
        raise SystemExit(2)
    source_dir = Path(sys.argv[1])
    output = Path(sys.argv[2])
    manifest_path = source_dir / "export-manifest.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8")) if manifest_path.exists() else {}
    exports = {}
    for workbook_path in sorted(source_dir.glob("*.xlsx")):
        exports[workbook_path.stem] = {"source_file": str(workbook_path), "sheets": read_workbook(workbook_path)}
    payload = {
        "item_id": manifest.get("item_id"),
        "captured_at": manifest.get("captured_at", datetime.now().isoformat(timespec="seconds")),
        "source_manifest": str(manifest_path) if manifest_path.exists() else None,
        "exports": exports,
        "errors": manifest.get("errors", []),
    }
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"OK {output} exports={len(exports)}")


if __name__ == "__main__":
    main()
