#!/usr/bin/env python3
"""创作者中心 作品列表.xlsx -> works-metrics.json（扁平 JSON，稳定英文 key）"""
import json
import sys
from datetime import datetime

from openpyxl import load_workbook

HEADER_MAP = {
    "作品名称": "title",
    "发布时间": "publish_time",
    "体裁": "format",
    "审核状态": "status",
    "播放量": "play_count",
    "完播率": "finish_rate",
    "5s完播率": "finish_rate_5s",
    "2s跳出率": "bounce_rate_2s",
    "封面点击率": "cover_ctr",
    "平均播放时长": "avg_watch_seconds",
    "点赞量": "like_count",
    "评论量": "comment_count",
 "分享量": "share_count",
    "收藏量": "collect_count",
    "主页访问量": "profile_visits",
    "粉丝增量": "follower_delta",
}

RATIO_KEYS = {"finish_rate", "finish_rate_5s", "bounce_rate_2s", "cover_ctr"}

def main():
    if len(sys.argv) < 2:
        print("用法: python3 parse_works_xlsx.py <xlsx> [输出.json]")
        sys.exit(2)
    xlsx_path = sys.argv[1]
    out_path = sys.argv[2] if len(sys.argv) > 2 else None
    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    ws = wb.active
    header = None
    works = []
    for row in ws.iter_rows(values_only=True):
        if header is None:
            if row and "作品名称" in row:
                header = list(row)
            continue
        if not row or not row[0]:
            continue
        record = {}
        for i, col in enumerate(header):
            key = HEADER_MAP.get(col)
            if key is None:
                continue
            value = row[i] if i < len(row) else None
            if isinstance(value, datetime):
                value = value.isoformat(sep=" ", timespec="minutes")
            if isinstance(value, str) and value.endswith("%"):
                value = str(value).rstrip("%")
            if key in RATIO_KEYS:
                value = round(float(value) * 100, 2) if value not in (None, "") else None
            if isinstance(value, str) and value.replace(".", "", 1).isdigit():
                value = float(value) if "." in value else int(value)
            record[key] = value
        if record.get("play_count") is None and record.get("publish_time") is None:
            continue
        works.append(record)
    payload = {
        "captured_at": datetime.now().isoformat(timespec="seconds"),
        "source_file": xlsx_path,
        "works": works,
        "work_count": len(works)
    }
    text = json.dumps(payload, ensure_ascii=False, indent=1)
    if out_path:
        with open(out_path, "w", encoding="utf-8") as f:
            f.write(text)
        print(f"OK {out_path} works={len(works)}")
    else:
        print(text)

if __name__ == "__main__":
    main()
